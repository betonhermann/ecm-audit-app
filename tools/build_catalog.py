"""
Erzeugt dist/js/catalog.js aus der Excel-Auditcheckliste.

Aufruf:  python3 tools/build_catalog.py <Auditcheckliste.xlsm>

Gelesen werden die (ausgeblendeten) Blätter ECM_Funktion_I bis _IV.
Spalte A = Nummer, B = Text, C-F = Ankreuzfelder (e., n.e., n.g., n.z.).
Zeilen mit Ankreuzfeldern sind bewertbare Fragen.
"""
import json, re, sys
import openpyxl

SHEETS = ["ECM_Funktion_I", "ECM_Funktion_II", "ECM_Funktion_III", "ECM_Funktion_IV"]
BLOCK_IDS = ("I", "I.I", "II", "III", "IV")

# CRSC-Zusatzfragen in ECM I (aus Modul1_Checkliste, ChecklisteErstellen)
CRSC_IDS = {"2.5", "2.6", "2.7", "2.8", "2.9", "2.10", "2.11", "2.12",
            "6.3", "6.3.1", "6.3.2", "7.5", "7.5.1", "7.5.2", "7.5.3", "7.5.4",
            "8.3", "8.3.1", "8.3.2", "9.7", "9.7.1", "9.7.2", "9.7.3"}


def clean(s):
    if s is None:
        return ""
    s = str(s).replace("_x000D_", "").replace("\r", "")
    s = re.sub(r"[ \t]+\n", "\n", s)
    return s.strip()


def reflow(text):
    """Harte Zeilenumbrüche aus der Excel-Zellbreite entfernen, echte Absätze und Listen behalten."""
    lines = text.split("\n")
    out = lines[0]
    for line in lines[1:]:
        stripped = line.strip()
        prev = out.rstrip()
        keep = (
            not stripped
            or not prev
            or prev[-1] in ":?.;!)"
            or re.match(r"^(\(|-|–|•|[a-z]\.\s|\d+\.\s)", stripped)
        )
        out = (out + "\n" + stripped) if keep else (prev + " " + stripped)
    return re.sub(r"\n{3,}", "\n\n", out).strip()


def parse(wb):
    blocks = []
    for sn in SHEETS:
        ws = wb[sn]
        block = None
        section = None
        sub_sel = None
        for r in range(1, ws.max_row + 1):
            a = clean(ws.cell(r, 1).value)
            b = clean(ws.cell(r, 2).value)
            if not a and not b:
                continue
            ratable = isinstance(ws.cell(r, 3).value, bool)
            is_section = ws.cell(r, 7).value == "Abweichungen"

            if a in BLOCK_IDS:
                block = {"id": a, "title": b, "desc": [], "sections": []}
                blocks.append(block)
                section = None
                continue

            if b.startswith("Nachweise / Bemerkungen") or (ratable and not b):
                # Nachweiszeile des Kapitels (I.I 10.2 hat in der Excel keinen Text)
                if section is not None:
                    section["evidence"] = True
                continue

            if is_section and a.endswith(".0"):
                a = a[:-2]  # II "6.0" ist ein normales Kapitel 6
            if is_section and "." in a and section is not None and block["id"] == "I.I":
                # Unterkapitel mit eigener Auswahl (I.I 11.2 / 11.3)
                sub_sel = a
                section["items"].append({"t": "sub", "id": a, "text": b, "sel": a})
                continue

            if is_section:
                section = {"id": a, "title": b, "evidence": False, "items": []}
                block["sections"].append(section)
                sub_sel = None
                continue

            if section is None:
                block["desc"].append({"id": a, "text": reflow(b)})
                continue

            if block["id"] == "I.I" and section["id"] == "11" and a == "11.1":
                sub_sel = "11.1"
                section["items"].append({"t": "sub", "id": a, "text": b, "sel": a})
                continue

            if block["id"] in ("I", "I.I"):
                sel = sub_sel or section["id"]
            else:
                sel = block["id"]

            if ratable:
                kind = "q"
            elif b.startswith("Hinweise für den Auditor"):
                kind = "hint"
                b = b.split("\n", 1)[1].strip() if "\n" in b else ""
            elif "CRSC-Lieferantenüberwachung" in b:
                kind = "crsc"
            else:
                kind = "text"

            item = {"t": kind, "id": a, "text": reflow(b), "sel": sel}
            if block["id"] == "I" and a in CRSC_IDS:
                item["crsc"] = True
            section["items"].append(item)
    return blocks


def main():
    src = sys.argv[1]
    wb = openpyxl.load_workbook(src)
    blocks = parse(wb)
    keys = set()
    n = 0
    for bl in blocks:
        for s in bl["sections"]:
            for it in s["items"]:
                if it["t"] == "q":
                    k = bl["id"] + "|" + it["id"]
                    if k in keys:
                        raise SystemExit("Doppelte Fragen-ID: " + k)
                    keys.add(k)
                    n += 1
    out = {"quelle": src.split("/")[-1], "fragen": n, "blocks": blocks}
    js = ("// Automatisch erzeugt mit tools/build_catalog.py, nicht von Hand bearbeiten.\n"
          "window.ECM_CATALOG = " + json.dumps(out, ensure_ascii=False, indent=1) + ";\n")
    open("dist/js/catalog.js", "w", encoding="utf-8").write(js)
    print(f"{n} Fragen in {len(blocks)} Blöcken geschrieben")


if __name__ == "__main__":
    main()
